"""Tests de l'API calculBA.  Lancer avec :  python -m pytest -q"""
import io
import os
import sqlite3
import tempfile

_TMP = tempfile.mkdtemp()
os.environ["CALCULBA_DB"] = os.path.join(_TMP, "test.db")

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from backend import settings  # noqa: E402
from backend.analysis import to_number  # noqa: E402
from backend.main import app  # noqa: E402

H = {settings.CSRF_HEADER: settings.CSRF_VALUE}


def make_client(email="ange@example.com"):
    c = TestClient(app)
    c.__enter__()
    r = c.post("/api/auth/register", headers=H, json={"full_name": "Ange Test", "email": email, "password": "motdepasse123"})
    assert r.status_code in (201, 409), r.text
    if r.status_code == 409:
        c.post("/api/auth/login", headers=H, json={"email": email, "password": "motdepasse123"})
    return c


@pytest.fixture(scope="module")
def client():
    return make_client()


def new_dataset(client, columns=("Echantillon", "Age_jours", "Resistance_MPa"), name="Essais"):
    r = client.post("/api/datasets", headers=H, json={"name": name, "domain": "Béton", "columns": list(columns)})
    assert r.status_code == 201, r.text
    return r.json()["id"]


# ------------------------------------------------------------------ nombres

@pytest.mark.parametrize("raw,expected", [
    ("22,5", 22.5), ("22.5", 22.5), (" 1 234,5 ", 1234.5), ("1.234,5", 1234.5), ("1,234.5", 1234.5),
    ("-3", -3.0), ("2e3", 2000.0), (7, 7.0), ("B01", None), ("", None), (None, None), ("12 kN", None), (True, None),
])
def test_to_number(raw, expected):
    assert to_number(raw) == expected


# --------------------------------------------------------------------- auth

def test_csrf_header_required(client):
    r = client.post("/api/datasets", json={"name": "x", "columns": ["a"]})
    assert r.status_code == 403


def test_birth_date_optional_and_not_exposed():
    c = make_client("sans.date@example.com")
    user = c.get("/api/auth/me").json()["user"]
    assert "birth_date" not in user


def test_register_rejects_implausible_birth_date():
    c = TestClient(app)
    r = c.post("/api/auth/register", headers=H, json={"full_name": "Vieux", "email": "v@example.com",
                                                      "password": "motdepasse123", "birth_date": "0001-01-01"})
    assert r.status_code == 422 and "date" in r.json()["detail"]


def test_wrong_password_and_throttle():
    c = TestClient(app)
    for _ in range(settings.LOGIN_MAX_FAILURES):
        assert c.post("/api/auth/login", headers=H, json={"email": "nobody@example.com", "password": "x" * 10}).status_code == 401
    assert c.post("/api/auth/login", headers=H, json={"email": "nobody@example.com", "password": "x" * 10}).status_code == 429


def test_expired_session_rejected():
    c = make_client("expire@example.com")
    assert c.get("/api/auth/me").status_code == 200
    from backend.db import connect
    db = connect()
    db.execute("UPDATE user_sessions SET expires_at = '2000-01-01 00:00:00' WHERE user_id = "
               "(SELECT id FROM users WHERE email = 'expire@example.com')")
    db.commit(); db.close()
    assert c.get("/api/auth/me").status_code == 401


def test_logout_invalidates_session():
    c = make_client("logout@example.com")
    c.post("/api/auth/logout", headers=H)
    assert c.get("/api/auth/me").status_code == 401


# ------------------------------------------------------------------ datasets

def test_isolation_between_users(client):
    did = new_dataset(client)
    other = make_client("autre@example.com")
    assert other.get(f"/api/datasets/{did}").status_code == 404
    assert other.get(f"/api/datasets/{did}/records").status_code == 404
    assert other.delete(f"/api/datasets/{did}", headers=H).status_code == 404


def test_duplicate_columns_rejected(client):
    r = client.post("/api/datasets", headers=H, json={"name": "x", "columns": ["a", "A"]})
    assert r.status_code == 422


def test_unknown_domain_rejected(client):
    r = client.post("/api/datasets", headers=H, json={"name": "x", "domain": "Inconnu", "columns": ["a"]})
    assert r.status_code == 422


def test_records_column_named_id_does_not_clash(client):
    did = new_dataset(client, columns=("id", "valeur"))
    rid = client.post(f"/api/datasets/{did}/records", headers=H, json={"values": {"id": "P-1", "valeur": "3"}}).json()["id"]
    rec = client.get(f"/api/datasets/{did}/records").json()["records"][0]
    assert rec["id"] == rid and rec["values"]["id"] == "P-1"


def test_bad_record_payloads(client):
    did = new_dataset(client)
    assert client.post(f"/api/datasets/{did}/records", headers=H, json={"values": [1]}).status_code == 422
    assert client.post(f"/api/datasets/{did}/records", headers=H, json={"values": {"Inconnue": 1}}).status_code == 400
    assert client.post(f"/api/datasets/{did}/records", headers=H, json={"values": {"Echantillon": "  "}}).status_code == 400


def test_edit_and_delete_record(client):
    did = new_dataset(client)
    rid = client.post(f"/api/datasets/{did}/records", headers=H, json={"values": {"Echantillon": "B1", "Age_jours": "7"}}).json()["id"]
    r = client.put(f"/api/datasets/{did}/records/{rid}", headers=H, json={"values": {"Echantillon": "B1", "Age_jours": "28"}})
    assert r.json()["values"]["Age_jours"] == "28"
    assert client.delete(f"/api/datasets/{did}/records/{rid}", headers=H).status_code == 200
    assert client.get(f"/api/datasets/{did}").json()["record_count"] == 0


# --------------------------------------------------------------------- stats

def test_type_detection_and_stats(client):
    did = new_dataset(client)
    for e, a, r in [("B01", "7", "22,5"), ("B02", "28", "31.8"), ("3", "7", ""), ("B04", "28", "33,1"), ("B05", "14", "27.9")]:
        client.post(f"/api/datasets/{did}/records", headers=H, json={"values": {"Echantillon": e, "Age_jours": a, "Resistance_MPa": r}})
    s = client.get(f"/api/datasets/{did}/stats").json()
    # Une seule valeur « 3 » parmi des identifiants ne rend pas la colonne numérique.
    assert s["numeric_columns"] == ["Age_jours", "Resistance_MPa"]
    res = s["columns"]["Resistance_MPa"]
    assert res["count"] == 4 and res["missing"] == 1 and res["invalid_rows"] == []
    assert res["mean"] == pytest.approx((22.5 + 31.8 + 33.1 + 27.9) / 4, abs=1e-4)
    # La valeur manquante n'est jamais comptée comme 0 dans l'histogramme.
    assert sum(res["histogram"]["counts"]) == 4 and res["min"] == 22.5

    g = client.get(f"/api/datasets/{did}/charts/groups", params={"by": "Age_jours", "value": "Resistance_MPa"}).json()
    assert [x["group"] for x in g["groups"]] == ["7", "14", "28"]
    sc = client.get(f"/api/datasets/{did}/charts/scatter", params={"x": "Age_jours", "y": "Resistance_MPa"}).json()
    assert sc["n"] == 4 and sc["fit"]["r2"] is not None
    se = client.get(f"/api/datasets/{did}/charts/series", params={"column": "Resistance_MPa"}).json()
    assert [p["x"] for p in se["points"]] == [1, 2, 4, 5]


def test_invalid_rows_are_located(client):
    # Les valeurs fautives arrivent par import (la saisie manuelle les refuse).
    csv = "Echantillon;R\nB1;20,5\nB2;rompu\nB3;22\nB4;\nB5;23\nB6;24\n".encode()
    did = upload(client, "rompu.csv", csv).json()["dataset"]["id"]
    s = client.get(f"/api/datasets/{did}/stats").json()["columns"]["R"]
    assert s["invalid"] == 1 and s["invalid_rows"] == [2]


def test_text_rejected_in_numeric_column_on_add_and_edit(client):
    did = new_dataset(client, columns=("Echantillon", "R"))
    ids = [client.post(f"/api/datasets/{did}/records", headers=H, json={"values": {"Echantillon": e, "R": r}}).json()["id"]
           for e, r in [("B1", "20,5"), ("B2", "21"), ("B3", "22,5")]]
    r = client.post(f"/api/datasets/{did}/records", headers=H, json={"values": {"Echantillon": "B4", "R": "abc"}})
    assert r.status_code == 400 and "non numérique" in r.json()["detail"]
    r = client.put(f"/api/datasets/{did}/records/{ids[0]}", headers=H, json={"values": {"Echantillon": "B1", "R": "abc"}})
    assert r.status_code == 400 and "R" in r.json()["detail"]
    # Texte libre toujours accepté dans une colonne texte, et champ vide accepté partout.
    assert client.put(f"/api/datasets/{did}/records/{ids[0]}", headers=H, json={"values": {"Echantillon": "Éprouvette 1", "R": ""}}).status_code == 200
    # Une colonne encore vide accepte n'importe quelle première valeur.
    did2 = new_dataset(client, columns=("Note",))
    assert client.post(f"/api/datasets/{did2}/records", headers=H, json={"values": {"Note": "abc"}}).status_code == 201


def test_single_value_std_is_none(client):
    did = new_dataset(client, columns=("v",))
    client.post(f"/api/datasets/{did}/records", headers=H, json={"values": {"v": "5"}})
    assert client.get(f"/api/datasets/{did}/stats").json()["columns"]["v"]["std"] is None


# -------------------------------------------------------------------- import

def upload(client, name, content, **form):
    return client.post("/api/datasets/import", headers=H, files={"file": (name, content)}, data=form)


def test_import_semicolon_csv_cp1252(client):
    csv = "Site;Profondeur_m;Teneur_eau_pct\nA;1,5;12,3\nB;2,0;15\n\n".encode("cp1252")
    r = upload(client, "géotech.csv", csv, domain="Sols et géotechnique")
    assert r.status_code == 201, r.text
    ds = r.json()["dataset"]
    assert ds["columns"] == ["Site", "Profondeur_m", "Teneur_eau_pct"] and r.json()["rows_imported"] == 2
    s = client.get(f"/api/datasets/{ds['id']}/stats").json()
    assert s["columns"]["Profondeur_m"]["mean"] == 1.75


def test_import_rejects_duplicate_and_unnamed_columns(client):
    assert "deux fois" in upload(client, "d.csv", b"a,a,b\n1,2,3\n").json()["detail"]
    assert "pas de nom" in upload(client, "e.csv", b"a,,b\n1,2,3\n").json()["detail"]
    # Une colonne vide sans en-tête (fréquent avec Excel) est simplement ignorée.
    r = upload(client, "f.csv", b"a,b,\n1,2,\n")
    assert r.status_code == 201 and r.json()["dataset"]["columns"] == ["a", "b"]


def test_import_xls_and_unknown(client):
    assert ".xlsx" in upload(client, "old.xls", b"x").json()["detail"]
    assert upload(client, "notes.pdf", b"x").status_code == 400


def test_import_and_export_xlsx(client):
    from openpyxl import Workbook, load_workbook
    wb = Workbook(); ws = wb.active
    ws.append(["Point", "Est_m", "Nord_m"]); ws.append(["P1", 100.5, 200]); ws.append(["P2", 101, 201.25])
    buf = io.BytesIO(); wb.save(buf)
    r = upload(client, "topo.xlsx", buf.getvalue())
    assert r.status_code == 201, r.text
    did = r.json()["dataset"]["id"]
    x = client.get(f"/api/datasets/{did}/export", params={"format": "xlsx"})
    wb2 = load_workbook(io.BytesIO(x.content))
    rows = list(wb2.active.values)
    assert rows[0] == ("Point", "Est_m", "Nord_m") and rows[1] == ("P1", 100.5, 200)


def test_csv_export_formats_and_formula_guard(client):
    did = new_dataset(client, columns=("Note", "Valeur"))
    client.post(f"/api/datasets/{did}/records", headers=H, json={"values": {"Note": "=SOMME(A1)", "Valeur": "2,5"}})
    fr = client.get(f"/api/datasets/{did}/export", params={"format": "csv-excel"}).content.decode("utf-8-sig")
    std = client.get(f"/api/datasets/{did}/export", params={"format": "csv"}).content.decode("utf-8-sig")
    assert fr.splitlines()[1] == "'=SOMME(A1);2,5"
    assert std.splitlines()[1] == "'=SOMME(A1),2.5"


# ----------------------------------------------------------------- migration

def test_migration_from_v3(tmp_path, monkeypatch, sqlite_only):
    from backend import db
    path = tmp_path / "v3.db"
    old = sqlite3.connect(path)
    old.executescript("""
        CREATE TABLE users (id INTEGER PRIMARY KEY AUTOINCREMENT, full_name TEXT NOT NULL, birth_date TEXT NOT NULL,
            email TEXT NOT NULL UNIQUE COLLATE NOCASE, password_hash TEXT NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE user_sessions (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, token_hash TEXT NOT NULL UNIQUE,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE datasets (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, name TEXT NOT NULL, domain TEXT NOT NULL,
            columns_json TEXT NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE records (id INTEGER PRIMARY KEY AUTOINCREMENT, dataset_id INTEGER NOT NULL, values_json TEXT NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP);
    """)
    import hashlib, secrets as s
    salt = s.token_bytes(16)
    legacy = salt.hex() + ":" + hashlib.pbkdf2_hmac("sha256", b"ancienmotdepasse", salt, 310000).hex()
    old.execute("INSERT INTO users (full_name, birth_date, email, password_hash) VALUES ('Ancien', '1999-01-01', 'old@example.com', ?)", (legacy,))
    old.execute("INSERT INTO user_sessions (user_id, token_hash) VALUES (1, 'abc')")
    old.execute("INSERT INTO datasets (user_id, name, domain, columns_json) VALUES (1, 'Vieux jeu', 'Béton', '[\"a\"]')")
    old.execute("INSERT INTO records (dataset_id, values_json) VALUES (1, '{\"a\": \"1\"}')")
    old.commit(); old.close()

    monkeypatch.setattr(settings, "DB_PATH", path)
    db.init_db()
    with TestClient(app) as c:
        r = c.post("/api/auth/login", headers=H, json={"email": "old@example.com", "password": "ancienmotdepasse"})
        assert r.status_code == 200, r.text
        ds = c.get("/api/datasets").json()
        assert ds[0]["name"] == "Vieux jeu" and ds[0]["record_count"] == 1
    conn = sqlite3.connect(path)
    assert conn.execute("SELECT password_hash FROM users").fetchone()[0].startswith("pbkdf2_sha256$600000$")
    assert conn.execute("SELECT COUNT(*) FROM user_sessions WHERE token_hash='abc'").fetchone()[0] == 0
    assert conn.execute("PRAGMA table_info(users)").fetchall()[2][3] == 0  # birth_date n'est plus NOT NULL


def test_api_docs_page_is_local_and_grouped():
    c = TestClient(app)
    r = c.get("/docs")
    assert r.status_code == 200 and "/static/vendor/swagger/swagger-ui-bundle.js" in r.text and "cdn." not in r.text
    spec = c.get("/openapi.json").json()
    assert [t["name"] for t in spec["tags"]] == ["Comptes", "Jeux de données", "Observations", "Analyse", "Service"]
    assert spec["paths"]["/api/datasets/{did}/stats"]["get"]["summary"] == "Statistiques descriptives"

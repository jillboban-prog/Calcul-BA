"""Lecture des fichiers, détection des types et calculs statistiques.

Principe : les valeurs sont conservées telles que saisies (« 22,5 », « B01 »,
« 007 »). La conversion en nombre se fait uniquement au moment de l'analyse,
avec une seule fonction (`to_number`) utilisée partout, pour que tableaux,
statistiques, graphiques et exports donnent toujours les mêmes chiffres.
"""
from __future__ import annotations

import csv
import io
import math
import re
from datetime import date, datetime, time
from typing import Any, Iterable

import numpy as np
import pandas as pd

from . import settings

# ----------------------------------------------------------------- nombres

_SPACES = re.compile(r"[\s  ']")  # espaces, insécables, apostrophe suisse
_NUMBER = re.compile(r"[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?")
# Caractères de contrôle (ex. Ctrl-Z de fin de fichier DOS) : invisibles, et un
# classeur .xlsx qui en contient est refusé par openpyxl. Tab, saut de ligne et
# retour chariot sont conservés.
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def to_number(value: Any) -> float | None:
    """Convertit une valeur saisie en nombre, ou None si ce n'en est pas un.

    Gère la virgule décimale (« 22,5 »), les séparateurs de milliers
    (« 1 234,5 », « 1.234,5 », « 1,234.5 ») et la notation scientifique.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, np.integer, np.floating)):
        f = float(value)
        return f if math.isfinite(f) else None
    s = _SPACES.sub("", str(value))
    if not s:
        return None
    if "," in s and "." in s:
        # Le séparateur le plus à droite est le séparateur décimal.
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif s.count(",") == 1:
        s = s.replace(",", ".")
    if not _NUMBER.fullmatch(s):
        return None
    f = float(s)
    return f if math.isfinite(f) else None


def is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not _CTRL.sub("", value).strip())


def clean_value(value: Any) -> Any:
    """Normalise une valeur avant stockage (sans la convertir)."""
    if value is None:
        return None
    if isinstance(value, (bool, np.bool_)):
        return "VRAI" if value else "FAUX"
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        f = float(value)
        if not math.isfinite(f):
            return None
        return int(f) if f.is_integer() and abs(f) < 1e15 else f
    if isinstance(value, int):
        return value
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat(sep=" ") if (value.hour or value.minute or value.second) else value.date().isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    s = _CTRL.sub("", str(value)).strip()
    if not s:
        return None
    return s[: settings.MAX_VALUE_LENGTH]


# ----------------------------------------------------------- import de fichiers

class ImportProblem(ValueError):
    """Erreur lisible renvoyée telle quelle à l'utilisateur."""


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        # Repli : le séparateur le plus fréquent sur la ligne d'en-tête.
        first = sample.splitlines()[0] if sample else ""
        return max(";,\t|", key=first.count)


def read_table(filename: str, raw: bytes) -> tuple[list[str], list[list[Any]]]:
    """Retourne (colonnes, lignes) à partir d'un CSV ou d'un classeur .xlsx."""
    name = filename.lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        text = _decode(raw)
        if not text.strip():
            raise ImportProblem("Le fichier est vide.")
        delimiter = _sniff_delimiter(text[:20_000])
        try:
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        except csv.Error as exc:  # cellule démesurée (> 131 072 caractères), etc.
            raise ImportProblem(
                "Ce fichier CSV ne peut pas être lu (cellule démesurée ou caractère invalide). "
                "Vérifiez-le, ou enregistrez-le à nouveau en .csv ou .xlsx."
            ) from exc
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            frame = pd.read_excel(io.BytesIO(raw), header=None, dtype=object, engine="openpyxl")
        except Exception as exc:  # fichier corrompu, protégé par mot de passe…
            raise ImportProblem("Ce classeur Excel ne peut pas être lu. Enregistrez-le à nouveau au format .xlsx.") from exc
        frame = frame.astype(object).where(pd.notna(frame), None)
        rows = frame.values.tolist()
    elif name.endswith(".xls"):
        raise ImportProblem("L'ancien format .xls n'est pas pris en charge. Dans Excel : Fichier › Enregistrer sous › Classeur Excel (.xlsx).")
    else:
        raise ImportProblem("Format non reconnu. Importez un fichier .csv ou .xlsx.")

    # Supprime les lignes entièrement vides (fréquentes en fin de feuille).
    rows = [r for r in rows if any(not is_empty(v) for v in r)]
    if not rows:
        raise ImportProblem("Le fichier ne contient aucune donnée.")

    header = [clean_value(v) for v in rows[0]]
    body = rows[1:]
    width = max(len(header), *(len(r) for r in body)) if body else len(header)
    header += [None] * (width - len(header))

    # Colonnes sans en-tête : acceptées seulement si elles sont entièrement vides.
    keep: list[int] = []
    for i, h in enumerate(header):
        column_empty = all(i >= len(r) or is_empty(r[i]) for r in body)
        if h is None:
            if column_empty:
                continue
            raise ImportProblem(f"La colonne n° {i + 1} contient des valeurs mais n'a pas de nom sur la première ligne.")
        keep.append(i)

    columns = [str(header[i]) for i in keep]
    validate_columns(columns)
    if not body:
        raise ImportProblem("Le fichier ne contient que la ligne d'en-tête.")
    if len(body) > settings.MAX_ROWS:
        raise ImportProblem(f"Le fichier contient {len(body):,} lignes ; la limite est de {settings.MAX_ROWS:,}.".replace(",", " "))

    data = [[clean_value(r[i]) if i < len(r) else None for i in keep] for r in body]
    return columns, data


def validate_columns(columns: list[str]) -> list[str]:
    cleaned = [str(c).strip() for c in columns]
    if not cleaned:
        raise ImportProblem("Ajoutez au moins une colonne.")
    if len(cleaned) > settings.MAX_COLUMNS:
        raise ImportProblem(f"{len(cleaned)} colonnes trouvées ; la limite est de {settings.MAX_COLUMNS}.")
    seen: dict[str, str] = {}
    for c in cleaned:
        if not c:
            raise ImportProblem("Chaque colonne doit avoir un nom.")
        if len(c) > settings.MAX_COLUMN_NAME_LENGTH:
            raise ImportProblem(f"Le nom de colonne « {c[:30]}… » dépasse {settings.MAX_COLUMN_NAME_LENGTH} caractères.")
        key = c.casefold()
        if key in seen:
            raise ImportProblem(f"La colonne « {c} » apparaît deux fois. Chaque nom de colonne doit être unique.")
        seen[key] = c
    return cleaned


# ------------------------------------------------------------- profils et stats

def _r(x: float | None, nd: int = 4, sig: bool = True) -> float | None:
    if x is None or not math.isfinite(x):
        return None
    # Petites valeurs (perméabilité 1e-8 m/s, déformations…) : arrondir à nd décimales
    # les écraserait à 0. En dessous de ce seuil on garde 6 chiffres significatifs.
    # Les valeurs courantes (>= 0,01 pour nd = 4) sont arrondies comme avant.
    if sig and nd >= 4 and x != 0 and abs(x) < 10 ** (2 - nd):
        return float(f"{x:.6g}")
    return float(round(x, nd))


def _is_padded_code(filled: list) -> bool:
    """Codes à largeur fixe avec zéros de tête (001, 002… 120) : des identifiants, pas des mesures."""
    strs = [str(v).strip() for v in filled]
    if not all(re.fullmatch(r"[0-9]+", t) for t in strs):
        return False
    return any(len(t) > 1 and t.startswith("0") for t in strs) and len({len(t) for t in strs}) == 1


def profile(columns: list[str], rows: list[dict]) -> dict:
    """Type de chaque colonne et valeurs numériques extraites (alignées sur les lignes)."""
    out = {}
    for col in columns:
        raw = [r.get(col) for r in rows]
        filled = [v for v in raw if not is_empty(v)]
        nums = [to_number(v) for v in raw]
        parsed = sum(1 for v in nums if v is not None)
        numeric = (bool(filled) and parsed / len(filled) >= settings.NUMERIC_THRESHOLD
                   and not _is_padded_code(filled))
        out[col] = {"numeric": numeric, "raw": raw, "nums": nums, "filled": len(filled), "parsed": parsed}
    return out


def histogram(values: np.ndarray) -> dict:
    if values.size == 0:
        return {"edges": [], "counts": []}
    lo, hi = float(values.min()), float(values.max())
    if lo == hi:
        return {"edges": [_r(lo - 0.5), _r(hi + 0.5)], "counts": [int(values.size)]}
    edges = np.histogram_bin_edges(values, bins="auto")
    n = int(min(30, max(5, len(edges) - 1)))
    counts, edges = np.histogram(values, bins=n)
    return {"edges": [_r(e, 6) for e in edges], "counts": [int(c) for c in counts]}


def describe(columns: list[str], rows: list[dict]) -> dict:
    prof = profile(columns, rows)
    total = len(rows)
    result: dict[str, dict] = {}
    for col in columns:
        p = prof[col]
        missing = total - p["filled"]
        if p["numeric"]:
            v = np.array([x for x in p["nums"] if x is not None], dtype=float)
            n = int(v.size)
            mean = float(v.mean()) if n else None
            std = float(v.std(ddof=1)) if n > 1 else None
            q1, med, q3 = (np.percentile(v, [25, 50, 75]).tolist() if n else (None, None, None))
            cv = (std / abs(mean) * 100) if (std is not None and mean) else None
            result[col] = {
                "type": "numeric",
                "count": n,
                "distinct": int(np.unique(v).size),
                "missing": missing,
                "invalid": p["filled"] - p["parsed"],
                # Numéros de ligne (à partir de 1, ordre de saisie) des valeurs non numériques,
                # pour que l'interface puisse y mener directement.
                "invalid_rows": [i + 1 for i, (raw, num) in enumerate(zip(p["raw"], p["nums"]))
                                 if num is None and not is_empty(raw)][:20],
                "mean": _r(mean), "std": _r(std), "cv": _r(cv, 2),
                "min": _r(float(v.min())) if n else None,
                "q1": _r(q1), "median": _r(med), "q3": _r(q3),
                "max": _r(float(v.max())) if n else None,
                "histogram": histogram(v),
            }
        else:
            values = [str(x).strip() for x in p["raw"] if not is_empty(x)]
            counts = pd.Series(values, dtype="string").value_counts() if values else pd.Series(dtype="int64")
            result[col] = {
                "type": "text",
                "count": len(values),
                "missing": missing,
                "distinct": int(counts.size),
                "top": [{"value": str(k), "count": int(c)} for k, c in counts.head(5).items()],
            }
    return {
        "total_rows": total,
        "numeric_columns": [c for c in columns if prof[c]["numeric"]],
        "text_columns": [c for c in columns if not prof[c]["numeric"]],
        "columns": result,
    }


def _thin(indices: list[int], limit: int) -> list[int]:
    """Échantillonnage régulier (déterministe) pour limiter le nombre de points tracés."""
    if len(indices) <= limit:
        return indices
    step = len(indices) / limit
    return [indices[int(i * step)] for i in range(limit)]


def series(rows: list[dict], column: str, limit: int = 1500) -> dict:
    nums = [to_number(r.get(column)) for r in rows]
    present = [i for i, v in enumerate(nums) if v is not None]
    kept = _thin(present, limit)
    return {
        "column": column,
        "points": [{"x": i + 1, "y": _r(nums[i], 6)} for i in kept],
        "plotted": len(kept),
        "available": len(present),
        "total_rows": len(rows),
    }


def scatter(rows: list[dict], x: str, y: str, limit: int = 2000) -> dict:
    pairs = [(to_number(r.get(x)), to_number(r.get(y))) for r in rows]
    pairs = [(a, b) for a, b in pairs if a is not None and b is not None]
    n = len(pairs)
    fit = None
    if n >= 3:
        xs = np.array([a for a, _ in pairs]); ys = np.array([b for _, b in pairs])
        if xs.std() > 0 and ys.std() > 0:
            slope, intercept = np.polyfit(xs, ys, 1)
            r = float(np.corrcoef(xs, ys)[0, 1])
            fit = {"slope": _r(float(slope), 6), "intercept": _r(float(intercept), 6), "r": _r(r, sig=False), "r2": _r(r * r, sig=False),
                   "x_min": _r(float(xs.min()), 6), "x_max": _r(float(xs.max()), 6)}
    kept = _thin(list(range(n)), limit)
    return {"x": x, "y": y, "n": n, "plotted": len(kept),
            "points": [{"x": _r(pairs[i][0], 6), "y": _r(pairs[i][1], 6)} for i in kept], "fit": fit}


def groups(rows: list[dict], by: str, value: str, limit: int = 30) -> dict:
    buckets: dict[str, list[float]] = {}
    for r in rows:
        v = to_number(r.get(value))
        if v is None:
            continue
        key = r.get(by)
        key = "(vide)" if is_empty(key) else str(key).strip()
        buckets.setdefault(key, []).append(v)
    ordered = sorted(buckets.items(), key=lambda kv: (-len(kv[1]), kv[0]))
    out = []
    for key, vals in ordered[:limit]:
        a = np.array(vals)
        out.append({"group": key, "n": int(a.size), "mean": _r(float(a.mean())),
                    "std": _r(float(a.std(ddof=1))) if a.size > 1 else None,
                    "min": _r(float(a.min())), "max": _r(float(a.max()))})
    # Quand tous les groupes tiennent dans le graphique : ordre naturel
    # (numérique pour des âges ou profondeurs, alphabétique sinon).
    if len(buckets) <= limit:
        if all(to_number(g["group"]) is not None for g in out):
            out.sort(key=lambda g: to_number(g["group"]))
        else:
            out.sort(key=lambda g: g["group"].casefold())
    return {"by": by, "value": value, "groups": out, "total_groups": len(buckets), "shown": len(out)}


# ------------------------------------------------------------------- export

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: Any) -> Any:
    """Empêche qu'un texte soit interprété comme une formule par Excel (injection CSV)."""
    if isinstance(value, str):
        value = _CTRL.sub("", value)   # aussi pour les données déjà enregistrées
        if value.startswith(_FORMULA_START) and to_number(value) is None:
            return "'" + value
    return value


def export_rows(columns: list[str], rows: list[dict], numeric: Iterable[str]) -> list[list[Any]]:
    numeric = set(numeric)
    out = []
    for r in rows:
        line = []
        for c in columns:
            v = r.get(c)
            if c in numeric:
                n = to_number(v)
                line.append(n if n is not None else safe_cell(v))
            else:
                line.append(safe_cell(v))
        out.append(line)
    return out


def _fmt_number(x: float, decimal_comma: bool) -> str:
    s = repr(int(x)) if float(x).is_integer() and abs(x) < 1e15 else repr(x)
    return s.replace(".", ",") if decimal_comma else s


def to_csv(columns: list[str], rows: list[list[Any]], excel_fr: bool) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";" if excel_fr else ",", lineterminator="\r\n")
    w.writerow(columns)
    for r in rows:
        w.writerow(["" if v is None else _fmt_number(v, excel_fr) if isinstance(v, (int, float)) else v for v in r])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(title: str, columns: list[str], rows: list[list[Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", title)[:31] or "Données"
    ws.append([_CTRL.sub("", c) for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(10, len(c) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------ lois de probabilité
#
# Ajustement par maximum de vraisemblance (scipy). Les lois à support positif
# (lognormale, Weibull, exponentielle) sont ajustées avec un seuil fixé à 0 :
# c'est l'usage en génie civil (résistances, durées, débits) et cela évite les
# ajustements à 3 paramètres instables (ex. décalage de -2 000 MPa).

DISTRIBUTIONS: dict[str, dict] = {
    "normal": {"label": "Normale", "positive": False,
               "use": "Mesures symétriques autour d'une moyenne (résistances, dimensions, teneurs)."},
    "lognormal": {"label": "Lognormale", "positive": True,
                  "use": "Valeurs positives asymétriques (perméabilités, résistances de sols, durées)."},
    "weibull": {"label": "Weibull", "positive": True,
                "use": "Ruptures et résistances de matériaux fragiles, durées de vie."},
    "gumbel": {"label": "Gumbel", "positive": False,
               "use": "Valeurs extrêmes : crues maximales annuelles, vents, charges maximales."},
    "exponential": {"label": "Exponentielle", "positive": True,
                    "use": "Temps entre événements (pannes, arrivées)."},
    "uniform": {"label": "Uniforme", "positive": False,
                "use": "Valeurs également probables entre un minimum et un maximum."},
}


class FitProblem(ValueError):
    """Erreur lisible renvoyée telle quelle à l'utilisateur."""


def _fit(name: str, x: np.ndarray):
    """Retourne (loi scipy figée, paramètres lisibles, nombre de paramètres libres)."""
    from scipy import stats as st

    if name == "normal":
        mu, sigma = st.norm.fit(x)
        return st.norm(mu, sigma), [("μ (moyenne)", mu), ("σ (écart-type)", sigma)], 2
    if name == "lognormal":
        s, _, scale = st.lognorm.fit(x, floc=0)
        return st.lognorm(s, 0, scale), [("μ du logarithme", math.log(scale)), ("σ du logarithme", s)], 2
    if name == "weibull":
        c, _, scale = st.weibull_min.fit(x, floc=0)
        return st.weibull_min(c, 0, scale), [("k (forme)", c), ("λ (échelle)", scale)], 2
    if name == "gumbel":
        loc, scale = st.gumbel_r.fit(x)
        return st.gumbel_r(loc, scale), [("μ (position)", loc), ("β (échelle)", scale)], 2
    if name == "exponential":
        _, scale = st.expon.fit(x, floc=0)
        return st.expon(0, scale), [("λ (taux)", 1 / scale), ("Moyenne 1/λ", scale)], 1
    if name == "uniform":
        loc, scale = st.uniform.fit(x)
        return st.uniform(loc, scale), [("a (minimum)", loc), ("b (maximum)", loc + scale)], 2
    raise FitProblem(f"Loi inconnue : {name}.")


def fit_distributions(values: list[float], selected: str = "normal") -> dict:
    """Ajuste toutes les lois applicables, les classe (AIC) et détaille la loi choisie."""
    from scipy import stats as st

    if selected not in DISTRIBUTIONS:
        raise FitProblem(f"Loi inconnue : {selected}.")
    x = np.asarray([v for v in values if v is not None and math.isfinite(v)], dtype=float)
    n = int(x.size)
    if n < 5:
        raise FitProblem("Il faut au moins 5 valeurs numériques pour ajuster une loi.")
    if float(x.std()) == 0.0:
        raise FitProblem("Toutes les valeurs sont identiques : aucune loi ne peut être ajustée.")
    positive = bool(np.all(x > 0))

    comparison = []
    frozen = {}
    for name, meta in DISTRIBUTIONS.items():
        entry = {"key": name, "label": meta["label"], "use": meta["use"]}
        if meta["positive"] and not positive:
            entry.update(available=False, reason="exige des valeurs strictement positives")
            comparison.append(entry)
            continue
        try:
            with np.errstate(all="ignore"):
                dist, params, k = _fit(name, x)
                loglik = float(np.sum(dist.logpdf(x)))
            if not math.isfinite(loglik) or not all(math.isfinite(v) for _, v in params):
                raise ValueError("ajustement non fini")
            ks = st.kstest(x, dist.cdf)
            frozen[name] = (dist, params)
            entry.update(available=True, aic=_r(2 * k - 2 * loglik, 2, sig=False),
                         ks_d=_r(float(ks.statistic), 4, sig=False), ks_p=_r(float(ks.pvalue), 4, sig=False))
        except Exception:
            entry.update(available=False, reason="ajustement impossible sur ces données")
        comparison.append(entry)

    ranked = sorted((c for c in comparison if c["available"]), key=lambda c: c["aic"])
    for i, c in enumerate(ranked, start=1):
        c["rank"] = i
        c["delta_aic"] = _r(c["aic"] - ranked[0]["aic"], 2, sig=False)
    comparison.sort(key=lambda c: (not c["available"], c.get("rank", 99)))

    if selected not in frozen:
        reason = next(c.get("reason", "") for c in comparison if c["key"] == selected)
        raise FitProblem(f"La loi {DISTRIBUTIONS[selected]['label'].lower()} ne peut pas être utilisée ici : {reason}.")
    dist, params = frozen[selected]

    lo, hi = float(x.min()), float(x.max())
    pad = (hi - lo) * 0.15
    grid = np.linspace(lo - pad, hi + pad, 200)
    if DISTRIBUTIONS[selected]["positive"]:
        grid = grid[grid > 0]
    with np.errstate(all="ignore"):
        pdf = dist.pdf(grid)
    counts, edges = np.histogram(x, bins=int(min(30, max(5, math.ceil(math.sqrt(n))))))
    density = counts / (n * np.diff(edges))
    q = {p: float(dist.ppf(p)) for p in (0.05, 0.5, 0.95)}
    emp = {p: float(np.percentile(x, p * 100)) for p in (0.05, 0.5, 0.95)}

    sel = next(c for c in comparison if c["key"] == selected)
    return {
        "n": n,
        "summary": {"mean": _r(float(x.mean())), "std": _r(float(x.std(ddof=1))), "min": _r(lo), "max": _r(hi)},
        "selected": {"key": selected, "label": DISTRIBUTIONS[selected]["label"], "use": DISTRIBUTIONS[selected]["use"],
                     "parameters": [{"name": nm, "value": _r(float(v), 6)} for nm, v in params],
                     "aic": sel["aic"], "rank": sel["rank"], "ks_d": sel["ks_d"], "ks_p": sel["ks_p"]},
        "best": ranked[0]["key"],
        "quantiles": [{"p": p, "model": _r(q[p]), "empirical": _r(emp[p])} for p in (0.05, 0.5, 0.95)],
        "comparison": comparison,
        "plot": {"edges": [_r(float(e), 6) for e in edges], "density": [_r(float(d), 8) for d in density],
                 "x": [_r(float(v), 6) for v in grid],
                 "pdf": [(_r(float(v), 8) if math.isfinite(v) else None) for v in pdf]},
        "small_sample": n < 30,
    }
